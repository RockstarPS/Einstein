import openpyxl
import re
import sys

def get_string_from_bytes(val):
    if isinstance(val,bytes):
        return val.decode('utf-8')
    else:
        return val

def get_bool_val(val):
    ret_val = False
    if type(val) is bool:
        ret_val = val
    else:
        if val.strip(" ") == "TRUE":
            ret_val = True
    return ret_val

def generate_code_comabs_tx_cfg(excel_name,user_name,gen_time):
    '''
    This method is used to generate the ComAbsMdl_TxCfg.c and ComAbsMdl_TxCfg.h file
    :param excel_name:Input excel sheet
    :param user_name:user name to be displayed for file generation
    :param gen_time:file generation time
    :return:None
    '''

    work_book = openpyxl.load_workbook(excel_name)
    tx_sheet = work_book["TxSignals"]
    tx_signal_config=[]
    tx_signal_enum_list=[]
    general_iso_support = False
    for x in range(2, tx_sheet.max_row + 1):
        if get_bool_val(tx_sheet.cell(row=x, column=4).value) is True:
            tx_sig_dict = {'SignalName': get_string_from_bytes(tx_sheet.cell(row=x, column=1).value.strip(" ")),
                           'VSignalName': get_string_from_bytes(tx_sheet.cell(row=x, column=2).value.strip(" ")),
                           'VSignalEnum': get_string_from_bytes(tx_sheet.cell(row=x, column=3).value.strip(" ")),
                           'Lengthinbits': int(tx_sheet.cell(row=x, column=5).value),
                           'MsgName': get_string_from_bytes(tx_sheet.cell(row=x, column=6).value.strip(" ")),
                           'MessageCbk': get_bool_val(tx_sheet.cell(row=x, column=7).value),
                           'TxConfirm': get_bool_val(tx_sheet.cell(row=x, column=8).value),
                           'ValueRetention': get_bool_val(tx_sheet.cell(row=x, column=9).value),
                           'InitTrigger': get_bool_val(tx_sheet.cell(row=x, column=10).value),
                           'DynSignal': get_bool_val(tx_sheet.cell(row=x, column=11).value),
                           'SignalIsoSupport': get_bool_val(tx_sheet.cell(row=x, column=12).value),
                           }
            if get_bool_val(tx_sheet.cell(row=x, column=12).value) is True:
                general_iso_support = True
            if tx_sig_dict["Lengthinbits"] % 8 == 0:
                tx_sig_dict["LengthInBytes"] = int(tx_sig_dict["Lengthinbits"] / 8)
            else:
                tx_sig_dict["LengthInBytes"] = int(tx_sig_dict["Lengthinbits"] / 8) + 1
            tx_signal_enum_list.append(tx_sheet.cell(row=x, column=3).value.strip())
            tx_signal_config.append(tx_sig_dict)

    # SignalName VSignalName VsignalEnumValue Lengthinbits MsgName MessageCbk TxConfirm ValueRetention DynSignal
    sigu8_cnt = 0
    sigu16_cnt = 0
    sigu32_cnt = 0
    siguarr_cnt = 0
    conf_count = 0
    siguarr_byte_len = 0
    set_sig = 0
    set_sigArr = 0
    confirm_msg_list = {}
    confirm_sig_list = []
    confirm_iso_sig_list = []
    set_sig_list = []
    set_sig_list_u8Arr = []
    sig_list = {}
    cnt = 0
    set_sig_list_props={}
    set_sig_list_u8Arr_props={}
    signal_space_max=0
    signal_enum_space_max=0
    msg_name_space_max=0
    for signals in tx_signal_config:
        # space calculation
        if len(signals["MsgName"]) >msg_name_space_max:
            msg_name_space_max = len(signals["MsgName"])
        if len(signals["SignalName"]) >signal_space_max:
            signal_space_max = len(signals["SignalName"])
        if len(signals["VSignalEnum"]) >signal_enum_space_max:
            signal_enum_space_max = len(signals["VSignalEnum"])

        # print signals["Lengthinbits"]
        sig_list[signals["SignalName"]] = cnt

        # print byte_length
        if signals["MsgName"] not in confirm_msg_list:
            confirm_msg_list[signals["MsgName"]] = [(signals["SignalName"])]
        else:
            confirm_msg_list[signals["MsgName"]].append(signals["SignalName"])

        if signals["SignalIsoSupport"] is False:
            if signals["LengthInBytes"] <= 4:
                set_sig += 1
                set_sig_list.append(signals["SignalName"])
                set_sig_list_props[signals["SignalName"]]=signals["VSignalEnum"]
            else:
                set_sigArr += 1
                set_sig_list_u8Arr.append(signals["SignalName"])
                set_sig_list_u8Arr_props[signals["SignalName"]] = signals["VSignalEnum"]

        if signals["Lengthinbits"] <= 8 :
            if signals["ValueRetention"] is True:
                sigu8_cnt += 1
        elif 8 < signals["Lengthinbits"] <= 16 :
            if signals["ValueRetention"] is True:
                sigu16_cnt += 1
        elif 16 < signals["Lengthinbits"] <= 32:
            if signals["ValueRetention"] is True:
                sigu32_cnt += 1
        else:
            if signals["ValueRetention"] is True:
                siguarr_cnt += 1
                siguarr_byte_len += signals["LengthInBytes"]

        if signals["TxConfirm"] is True:
            conf_count += 1
            confirm_sig_list.append(signals["SignalName"])
            if signals['SignalIsoSupport'] is True:
                confirm_iso_sig_list.append(True)
            else:
                confirm_iso_sig_list.append(False)
        cnt += 1
    # *    Auto generated code in """ + username + """ at """ + str(current_time) + """
    # print confirm_msg_list
    file_cfg_h = open("ComAbsMdl_TxCfg.h", "w")
    stdout_fileno = sys.stdout
    sys.stdout = file_cfg_h
    print('''/********************************************************************************************************************
* C O P Y R I G H T                                                                                                 *
*********************************************************************************************************************
* Copyright (c) 2019 by Visteon Corporation.       All rights reserved.                                             *
*                                                                                                                   *
* NOTICE: This Is an unpublished work Of authorship, which contains trade secrets.                                  *
* Visteon Corporation owns all rights to this work And intends to maintain it in confidence to preserve             *
* its trade secret status. Visteon Corporation reserves the right, under the copyright laws of the United States    *
* Or those of any other country that may have jurisdiction, to protect this work as an unpublished work,            *
* in the event of an inadvertent Or deliberate unauthorized publication. Visteon Corporation also reserves its      *
* rights under all copyright laws to protect this work as a published work, when appropriate.                       *
* Those having access to this work may Not copy it, use it, modify it, Or disclose the information contained in it  *
* without the written authorization of Visteon Corporation.                                                         *
********************************************************************************************************************/
/*!*****************************************************************************************************************
*    @file ComAbsMdl_TxCfg.h
*    @ingroup ComAbsMdl
*    @brief This file contains the configuration of the ComAbsMdl component.
*    Auto generated code in ''' + user_name + ''' at ''' + str(gen_time) + '''
********************************************************************************************************************/
#ifndef COMABSMDL_TX_CFG_H_
#define COMABSMDL_TX_CFG_H_

/********************************************************************************************************************
*  PUBLIC MACRO DEFINITIONS                                                                                         *
********************************************************************************************************************/''')
    if general_iso_support is True:
        print('#define COMABSMDL_TXCAN_ISO_SIGNAL_SUPPORTED')
    if sigu8_cnt > 0:
        print('#define COMABSMDL_TXCAN_NUM_OF_U8_BUFFERS		 ' + str(sigu8_cnt) + 'U')
    else:
        print('#define COMABSMDL_TXCAN_NUM_OF_U8_BUFFERS		 1U')

    if sigu16_cnt > 0:
        print('#define COMABSMDL_TXCAN_NUM_OF_U16_BUFFERS		 ' + str(sigu16_cnt) + 'U')
    else:
        print('#define COMABSMDL_TXCAN_NUM_OF_U16_BUFFERS		 1U')

    if sigu32_cnt > 0:
        print('#define COMABSMDL_TXCAN_NUM_OF_U32_BUFFERS		 ' + str(sigu32_cnt) + 'U')
    else:
        print('#define COMABSMDL_TXCAN_NUM_OF_U32_BUFFERS		 1U')

    if siguarr_byte_len > 0:
        print('#define COMABSMDL_TXCAN_NUM_OF_U8ARR_BUFFERS		 ' + str(siguarr_byte_len) + 'U')
    else:
        print('#define COMABSMDL_TXCAN_NUM_OF_U8ARR_BUFFERS		 1U')

    if conf_count > 0:
        print('#define COMABSMDL_TXCAN_NUM_OF_CONF_STATUS_BYTES		 ' + str(int(len(tx_signal_config) / 8) + 1) + 'U')
    else:
        print('#define COMABSMDL_TXCAN_NUM_OF_CONF_STATUS_BYTES		 1U')

    if siguarr_cnt > 0:
        print('#define COMABSMDL_TXCAN_NUM_OF_U8ARR_SIGNAL		 ' + str(siguarr_cnt) + 'U')
    else:
        print('#define COMABSMDL_TXCAN_NUM_OF_U8ARR_SIGNAL		 1U')

    if set_sig > 0:
        print('#define CANABSMDL_TXCAN_NUM_OF_SET_SIG_FUNCT		 ' + str(set_sig) + 'U')
    else:
        print('#define CANABSMDL_TXCAN_NUM_OF_SET_SIG_FUNCT		 1U')
    print('#define COMABSMDL_TXCAN_NO_SET_SIG_FUNCT		     CANABSMDL_TXCAN_NUM_OF_SET_SIG_FUNCT')

    if set_sigArr > 0:
        print('#define CANABSMDL_TXCAN_NUM_OF_SET_SIG_U8ARR_FUNCT		 ' + str(set_sigArr) + 'U')
    else:
        print('#define CANABSMDL_TXCAN_NUM_OF_SET_SIG_U8ARR_FUNCT		 1U')
    print('#define COMABSMDL_TXCAN_NO_SET_SIG_U8ARR_FUNCT		     CANABSMDL_TXCAN_NUM_OF_SET_SIG_U8ARR_FUNCT')

    print('''
/********************************************************************************************************************
*  PUBLIC TYPE DEFINITIONS                                                                                         *
********************************************************************************************************************/
typedef Std_ReturnType (*ComAbsMdl_TxCanSetSignalFunctPtrType)(uint8 * pSignalValue);
typedef Std_ReturnType (*ComAbsMdl_TxCanSetSignalU8ArrFunctPtrType)(uint8 * pSignalValue,uint32 SignalLength);

#define COMABSMDL_TXCAN_SIG_STATUS_NONE        0x0U
#define COMABSMDL_TXCAN_KEEP_ALIVE_BIT_ENABLED 0x1U
#define COMABSMDL_TXCAN_CONFIRMATION_ENABLED   0x2U
#define COMABSMDL_TXCAN_ISO_SUPPORTED          0x4U
typedef uint8 ComAbsMdl_TxCanSigalPropsType;

typedef struct
{
    uint16    SignalLengthInByte;
    uint16    SignalBufferStartIdx;
    uint16	  TxSetSignalIdx;
    ComAbsMdl_TxCanSigalPropsType TxSigalProps;
}ComAbsMdl_TxCanSignalConfigType;

/********************************************************************************************************************
*  PUBLIC CONSTANT AND VARIABLE DECLARTAION                                                                         *
********************************************************************************************************************/
extern const ComAbsMdl_TxCanSignalConfigType ComAbsMdl_TxCanSignalConfig[eEndOfTxSignal];
extern const ComAbsMdl_TxCanSetSignalFunctPtrType ComAbsMdl_TxCanSetSignalFunctPtrConfig[CANABSMDL_TXCAN_NUM_OF_SET_SIG_FUNCT];
extern const ComAbsMdl_TxCanSetSignalU8ArrFunctPtrType ComAbsMdl_TxCanSetSignalU8ArrFunctPtrConfig[CANABSMDL_TXCAN_NUM_OF_SET_SIG_U8ARR_FUNCT];

/********************************************************************************************************************
*  PUBLIC FUNCTION PROTOTYPE                                                                         			    *
********************************************************************************************************************/
''')
    for items in set_sig_list:
        print('extern Std_ReturnType ComAbsMdl_TxCanSetSig_' + items + '(uint8 * pSignalValue);')
    for items in set_sig_list_u8Arr:
        print('extern Std_ReturnType ComAbsMdl_TxCanSetSig_' + items + '(uint8 * pSignalValue,uint32 SignalLength);')
    print('''
#ifdef COMABSMDL_TXCAN_ISO_SIGNAL_SUPPORTED
/*!******************************************************************************************************************
*   @brief   This method is used to transmit the Iso signal value to the COM.
*   @param[in] SignalId TX signal id
*   @param[in] pTxSignalData pointer to the data to be updated in COM
*   @param[in] SignalLength length of the data to be udpated
*   @return Std_ReturnType status of the request
*   E_OK status accepted
*   E_NOT_OK status not accepted
********************************************************************************************************************/
extern Std_ReturnType CComAbsMdl_TxCan_SetIsoSignal(tTxSignalId SignalId,uint8 *pTxSignalData,uint32 SignalLength);

/*!******************************************************************************************************************
*   @brief   This method is called from ComAbsMdl_Tx on successful transmission of the signal CanTp signal.
*   Based on the last transmission request the corresponding ISO signal confirmation to be udpated.
********************************************************************************************************************/
extern void CComAbsMdl_TxCan_SetIsoSignalConfirmation(tTxSignalId SignalId);
#endif
''')

    uarr_list_str=[]
    tmp_cnt =0
    print ('''/********************************************************************************************************************
*  Below are the set of signals that can be send via CComAbsMdl_TxCan_SetSignalValue API               			    *
********************************************************************************************************************/''')
    print("/* SignalEnum SignalEnumValue SignalNameInDbc MessageName")
    signal_space_max += 4
    signal_enum_space_max += 4
    msg_name_space_max += 4
    for signals in tx_signal_config:
        tmp_str="   "+signals["VSignalEnum"]+" "*(signal_enum_space_max-len(signals["VSignalEnum"]))+\
                                                  str(tmp_cnt)+" "*(6-len(str(tmp_cnt)))+signals["SignalName"]+\
                                                  " "*(signal_space_max-len(signals["SignalName"]))+signals["MsgName"]
        if signals["LengthInBytes"] <=4:
            print (tmp_str)
        else:
            uarr_list_str.append(tmp_str)
        tmp_cnt+=1
    print ('*/')
    print('''/********************************************************************************************************************
*  Below are the set of signals that can be send via CComAbsMdl_TxCan_SetSignalValueU8Arr API          			    *
********************************************************************************************************************/''')
    print("/* SignalEnum SignalEnumValue SignalNameInDbc MessageName")
    for val in uarr_list_str:
        print (val)
    print('*/')

    print("""
#endif /* COMABSMDL_TX_CFG_H_ */\n
/*******************************************************************************************************************
*                                 END OF FILE                                                                      *
********************************************************************************************************************/\n\n""")
    sys.stdout.close()
    sys.stdout = stdout_fileno
    file_cfg_h.close()

    file_cfg_h = open("ComAbsMdl_TxCfg.c", "w")
    stdout_fileno = sys.stdout
    sys.stdout = file_cfg_h

    print('''/********************************************************************************************************************
* C O P Y R I G H T                                                                                                 *
*********************************************************************************************************************
* Copyright (c) 2019 by Visteon Corporation.       All rights reserved.                                             *
*                                                                                                                   *
* NOTICE: This Is an unpublished work Of authorship, which contains trade secrets.                                  *
* Visteon Corporation owns all rights to this work And intends to maintain it in confidence to preserve             *
* its trade secret status. Visteon Corporation reserves the right, under the copyright laws of the United States    *
* Or those of any other country that may have jurisdiction, to protect this work as an unpublished work,            *
* in the event of an inadvertent Or deliberate unauthorized publication. Visteon Corporation also reserves its      *
* rights under all copyright laws to protect this work as a published work, when appropriate.                       *
* Those having access to this work may Not copy it, use it, modify it, Or disclose the information contained in it  *
* without the written authorization of Visteon Corporation.                                                         *
********************************************************************************************************************/
/*!*****************************************************************************************************************
*    @file ComAbsMdl_TxCfg.c
*    @ingroup ComAbsMdl
*    @brief This file contains the configuration of the ComAbsMdl component.
*    @Note : Auto Generated code .Do Not hand modify this file
*    Auto generated code in ''' + user_name + ''' at ''' + str(gen_time) + '''
********************************************************************************************************************/
#include  "ComAbsMdl.h"

/********************************************************************************************************************
*  PUBLIC CONSTANT And VARIABLE DEFINITION                                                                          *
********************************************************************************************************************/
const ComAbsMdl_TxCanSetSignalFunctPtrType ComAbsMdl_TxCanSetSignalFunctPtrConfig[CANABSMDL_TXCAN_NUM_OF_SET_SIG_FUNCT]=
{''')
    if len(set_sig_list) > 0:
        cnt = 0
        for items in set_sig_list:
            if cnt == len(set_sig_list) - 1:
                print('    &ComAbsMdl_TxCanSetSig_' + items+ '  /* '+str(cnt)+" - "+set_sig_list_props[items]+' */')
            else:
                print('    &ComAbsMdl_TxCanSetSig_' + items + ','+'  /* '+str(cnt)+" - "+set_sig_list_props[items]+' */')
            cnt += 1
    else:
        print('    NULL_PTR')
    print('''};

const ComAbsMdl_TxCanSetSignalU8ArrFunctPtrType ComAbsMdl_TxCanSetSignalU8ArrFunctPtrConfig[CANABSMDL_TXCAN_NUM_OF_SET_SIG_U8ARR_FUNCT]=
{''')
    if len(set_sig_list_u8Arr) > 0:
        cnt = 0
        for items in set_sig_list_u8Arr:
            if cnt == len(set_sig_list_u8Arr) - 1:
                print('    &ComAbsMdl_TxCanSetSig_' + items+ '  /* '+str(cnt)+" - "+set_sig_list_u8Arr_props[items]+' */')
            else:
                print('    &ComAbsMdl_TxCanSetSig_' + items + ','+ '  /* '+str(cnt)+" - "+set_sig_list_u8Arr_props[items]+' */')
            cnt += 1
    else:
        print('    NULL_PTR')
    print('''};

const ComAbsMdl_TxCanSignalConfigType ComAbsMdl_TxCanSignalConfig[eEndOfTxSignal]=
{''')
    l_u8 = 0
    l_u16 = 0
    l_u32 = 0
    l_u64 = 0
    cnt = 0
    set_id = 0
    set_id_arr = 0
    # SignalName VSignalName VSignalEnum Lengthinbits MsgName MessageCbk TxConfirm ValueRetention DynSignal
    '''{
       1U,
       COMABSMDL_TXCAN_NUM_OF_U8_BUFFERS,
       0U,
       COMABSMDL_TXCAN_CONFIRMATION_ENABLED 
    },'''
    digit_space_len=8
    byte_pos_space = len("COMABSMDL_TXCAN_NUM_OF_U8ARR_BUFFERS")+4
    for signals in tx_signal_config:
        byte_length = signals["LengthInBytes"]
        temp_str = '    { '
        if signals["Lengthinbits"] <= 8:
            if signals["ValueRetention"] is True:
                temp_str += str(byte_length) + 'U'+" "*(digit_space_len-(len(str(byte_length))+1)) +', '+ str(l_u8) +\
                                                               'U'+" "*(byte_pos_space-(len(str(l_u8))+1))
                l_u8 += 1
            else:
                temp_str += str(byte_length) + 'U' + " " * (digit_space_len - (len(str(byte_length)) + 1)) + ', ' + \
                            "COMABSMDL_TXCAN_NUM_OF_U8_BUFFERS" + " " * (byte_pos_space - (len("COMABSMDL_TXCAN_NUM_OF_U8_BUFFERS") ))
        elif 8 < signals["Lengthinbits"] <= 16:
            if signals["ValueRetention"] is True:
                temp_str += str(byte_length) + 'U' + " " * (digit_space_len - (len(str(byte_length)) + 1)) + ', ' + str(
                    l_u16) + 'U' + " " * (byte_pos_space - (len(str(l_u16)) + 1))
                l_u16 += 1
            else:
                temp_str += str(byte_length) + 'U' + " " * (digit_space_len - (len(str(byte_length)) + 1)) + ', ' +\
                            "COMABSMDL_TXCAN_NUM_OF_U16_BUFFERS"  + " " * (byte_pos_space - (len("COMABSMDL_TXCAN_NUM_OF_U16_BUFFERS") ))
        elif 16 < signals["Lengthinbits"] <= 32:
            if signals["ValueRetention"] is True:
                temp_str += str(byte_length) + 'U' + " " * (digit_space_len - (len(str(byte_length)) + 1)) + ', ' + str(
                    l_u32) + 'U' + " " * (byte_pos_space - (len(str(l_u32)) + 1))
                l_u32 += 1
            else:
                temp_str += str(byte_length) + 'U' + " " * (digit_space_len - (len(str(byte_length)) + 1)) + ', ' + \
                            "COMABSMDL_TXCAN_NUM_OF_U32_BUFFERS" + " " * (byte_pos_space - (len("COMABSMDL_TXCAN_NUM_OF_U32_BUFFERS") ))
        else:
            if signals["ValueRetention"] is True:
                temp_str += str(byte_length) + 'U' + " " * (digit_space_len - (len(str(byte_length)) + 1)) + ', ' + str(
                    l_u64) + 'U' + " " * (byte_pos_space - (len(str(l_u64)) + 1))
                l_u64 += byte_length
            else:
                temp_str += str(byte_length) + 'U' + " " * (digit_space_len - (len(str(byte_length)) + 1)) + ', ' + \
                            "COMABSMDL_TXCAN_NUM_OF_U8ARR_BUFFERS" + " " * (byte_pos_space - (len("COMABSMDL_TXCAN_NUM_OF_U8ARR_BUFFERS") ))

        if byte_length <= 4:
            temp_str += ', ' + str(set_id) + 'U'+" " * (digit_space_len - (len(str(set_id)) + 1))
            set_id += 1
        else:
            temp_str += ', ' + str(set_id_arr) + 'U'+" " * (digit_space_len - (len(str(set_id_arr)) + 1))
            set_id_arr += 1
        temp_flag_str = ', ( COMABSMDL_TXCAN_CONFIRMATION_ENABLED'
        if signals["ValueRetention"] is True:
            temp_flag_str += ' | COMABSMDL_TXCAN_KEEP_ALIVE_BIT_ENABLED'
        if signals['SignalIsoSupport'] is True:
            temp_flag_str += ' | COMABSMDL_TXCAN_ISO_SUPPORTED'
        space_len = len('( COMABSMDL_TXCAN_CONFIRMATION_ENABLED | COMABSMDL_TXCAN_KEEP_ALIVE_BIT_ENABLED | COMABSMDL_TXCAN_ISO_SUPPORTED )')
        temp_flag_str += ' ) ' + " " * (space_len - len(temp_flag_str) + 1)
        temp_str += temp_flag_str

        if cnt == len(tx_signal_config) - 1:
            temp_str += '} /* ' + str(cnt) + ' ' + signals["SignalName"] + " " + signals["VSignalEnum"] + ' */'
        else:
            temp_str += '}, /* ' + str(cnt) + ' ' + signals["SignalName"] + " " + signals["VSignalEnum"] + ' */'
        print(temp_str)
        cnt += 1
    print('''};
/********************************************************************************************************************
*  PUBLIC FUNCTION DEFINITION                                                                         	         	*
********************************************************************************************************************/
''')
    # confirmation
    cnt = 0
    for items in confirm_sig_list:
        print('''/********************************************************************************************************************
*   @brief   This method is called from com on successful transmission of the signal ''' + items + ''' 
********************************************************************************************************************/''')
        print('''void CComAbsMdl_OnDataTx_''' + items + '(void)')
        print('{')
        # print items
        sig_con = tx_signal_config[sig_list[items]]
        # sig_con["MsgName"]
        if sig_con["MsgName"] in confirm_msg_list:
            for sig in confirm_msg_list[sig_con["MsgName"]]:
                # print sig,confirm_msg_list
                sig_con_temp = tx_signal_config[sig_list[sig]]
                if sig_con_temp['SignalIsoSupport'] is True:
                    print ('    CComAbsMdl_TxCan_SetIsoSignalConfirmation(' + sig_con_temp["VSignalEnum"] + ');')
                else:
                    print('    ComAbsMdl_TxCan_SetCanSignalTxConfirmation(' + sig_con_temp["VSignalEnum"] + ');')
        print('}\n')
        cnt+=1
    # write
    for signals in tx_signal_config:
        if signals["Lengthinbits"] % 8 == 0:
            byte_length = int(signals["Lengthinbits"] / 8)
        else:
            byte_length = int(signals["Lengthinbits"] / 8) + 1
        if (signals["SignalIsoSupport"] is False):
            if (byte_length <= 4):
                print('''/********************************************************************************************************************
*   @brief   This method Is used to send the signal ''' + signals["SignalName"] + ''' to Com 
*   @param[in] SignalValue pointer to signal data
********************************************************************************************************************/
Std_ReturnType ComAbsMdl_TxCanSetSig_''' + signals["SignalName"] + '''(uint8 * pSignalValue)
{
    Std_ReturnType RetStatus;''')
                if byte_length == 1:
                    print('''    uint8   SignalValue = 0;''')
                    print('''    SignalValue = *((uint8 *)pSignalValue);''')
                if byte_length == 2:
                    print('''    uint16   SignalValue = 0;''')
                    print('''    SignalValue = *((uint16 *)pSignalValue);''')
                if 2 < byte_length <= 4:
                    print('''    uint32   SignalValue = 0;''')
                    print('''    SignalValue = *((uint32 *)pSignalValue);''')

                print('''    RetStatus = Rte_Write_pp_SR_Tx''' + signals["SignalName"] + "_" + signals[
                    "SignalName"] + '(SignalValue);')
                print('''    return RetStatus;
}\n''')
            else:
                print('''/********************************************************************************************************************
*   @brief   This method Is used to send the signal ''' + signals["SignalName"] + ''' to Com 
*   @param[in] SignalValue pointer to signal data
*   @param[in] SignalLength signal length
********************************************************************************************************************/
Std_ReturnType ComAbsMdl_TxCanSetSig_''' + signals["SignalName"] + '''(uint8 * pSignalValue,uint32 SignalLength)
{
    Std_ReturnType RetStatus;''')
                if signals["DynSignal"] is True:
                    print('''    RetStatus = Rte_Send_pp_SR_Tx''' + signals["SignalName"] + "_" + signals[
                        "SignalName"] + '(pSignalValue,SignalLength);')
                else:
                    print('''    RetStatus = Rte_Write_pp_SR_Tx''' + signals["SignalName"] + "_" + signals[
                        "SignalName"] + '(pSignalValue);')
                    print ('    COMABSMDL_DUMMY_STATEMENT(SignalLength);')
                print('''    return RetStatus;
}\n''')

    print('''/*******************************************************************************************************************
*                                 END OF FILE                                                                      *
********************************************************************************************************************/\n\n''')

    sys.stdout.close()
    sys.stdout = stdout_fileno
    file_cfg_h.close()


if __name__ == '__main__':
    import datetime
    import getpass
    excel_name = "ford_pdc (2).xlsx"
    current_time = datetime.datetime.now()
    username = getpass.getuser()
    generate_code_comabs_tx_cfg(excel_name, username,current_time)