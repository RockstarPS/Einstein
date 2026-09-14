import openpyxl
import sys
import datetime
import getpass

# ==================== CONFIGURATION ====================
MODULE_NAME = "ComAbsMdl"  # <<< Change this to ComAbsMdlSafe or any module
HEADER_FILE = f"{MODULE_NAME}_RxCfg.h"
SOURCE_FILE = f"{MODULE_NAME}_RxCfg.c"
PREFIX = MODULE_NAME.upper()


def get_string_from_bytes(val):
    if isinstance(val, bytes):
        return val.decode('utf-8').strip()
    else:
        return str(val).strip()


def get_bool_val(val):
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    return str(val).strip().upper() == "TRUE"


def get_ctype_for_bits(nbits):
    """Returns appropriate C type for a given number of bits"""
    if nbits <= 8:
        return "uint8"
    elif nbits <= 16:
        return "uint16"
    elif nbits <= 32:
        return "uint32"
    else:
        return "uint8"  # For consistency in scalar context; arrays use uint8* anyway


def generate_code_comabs_rx_cfg(excel_name, user_name, gen_time):
    try:
        workbook = openpyxl.load_workbook(excel_name, data_only=True)
    except Exception as e:
        print(f"Error opening Excel file: {e}")
        sys.exit(1)

    try:
        rx_sheet = workbook["RxSignals"]
    except KeyError:
        print("Error: Sheet 'RxSignals' not found in workbook.")
        sys.exit(1)

    # Lists and maps
    rx_signal_config = []
    rx_signal_enum_list = []
    rx_signal_qf_invalid_list = {}
    general_iso_support = False

    # Parse all enabled signals
    for row in range(2, rx_sheet.max_row + 1):
        if not get_bool_val(rx_sheet.cell(row=row, column=4).value):  # Enabled?
            continue
        try:
            enum_val = get_string_from_bytes(rx_sheet.cell(row=row, column=3).value)
            if not enum_val:
                print(f"Warning: Empty VSignalEnum at row {row}, skipping.")
                continue
            if enum_val in rx_signal_enum_list:
                raise ValueError(f"Duplicate VSignalEnum: {enum_val}")

            signal_dict = {
                'SignalName': get_string_from_bytes(rx_sheet.cell(row=row, column=1).value),
                'VSignalName': get_string_from_bytes(rx_sheet.cell(row=row, column=2).value),
                'VSignalEnum': enum_val,
                'Lengthinbits': int(rx_sheet.cell(row=row, column=5).value or 0),
                'TimeoutValue': int(rx_sheet.cell(row=row, column=6).value or 0),
                'MsgName': get_string_from_bytes(rx_sheet.cell(row=row, column=7).value),
                'DynSignalSupport': get_bool_val(rx_sheet.cell(row=row, column=8).value),
                'MessageCbk': get_bool_val(rx_sheet.cell(row=row, column=9).value),
                'SignalInvalidSupport': get_bool_val(rx_sheet.cell(row=row, column=10).value),
                'SignalInvalidVal': rx_sheet.cell(row=row, column=11).value,
                'SignalInvalidCounter': int(rx_sheet.cell(row=row, column=12).value or 0),
                'SignalUBSupport': get_bool_val(rx_sheet.cell(row=row, column=13).value),
                'SignalQfSupport': get_bool_val(rx_sheet.cell(row=row, column=14).value),
                'SignalQfSignalValue': get_string_from_bytes(rx_sheet.cell(row=row, column=15).value),
                'SignalInitTrigger': rx_sheet.cell(row=row, column=16).value,
                'ValueRetention': get_bool_val(rx_sheet.cell(row=row, column=17).value),
                'SignalIsoSupport': get_bool_val(rx_sheet.cell(row=row, column=18).value)
            }

            if signal_dict['SignalIsoSupport']:
                general_iso_support = True

            # Compute length in bytes
            bits = signal_dict['Lengthinbits']
            signal_dict['LengthInBytes'] = (bits + 7) // 8

            rx_signal_config.append(signal_dict)
            rx_signal_enum_list.append(enum_val)

            if signal_dict['SignalQfSupport'] and signal_dict['SignalQfSignalValue']:
                qf_enum = get_string_from_bytes(signal_dict['SignalQfSignalValue'])
                if qf_enum not in rx_signal_qf_invalid_list:
                    rx_signal_qf_invalid_list[qf_enum] = {
                        'SignalName': signal_dict['SignalName'],
                        'SignalInvalidVal': signal_dict['SignalInvalidVal'],
                        'SignalInvalidCounter': signal_dict['SignalInvalidCounter'],
                        'VSignalEnum': qf_enum
                    }
        except Exception as e:
            print(f"Error parsing row {row}: {e}")
            sys.exit(1)

    # Analyze signal types
    sigu8_cnt = sigu8_retention_cnt = 0
    sigu16_cnt = sigu16_retention_cnt = 0
    sigu32_cnt = sigu32_retention_cnt = 0
    sigu64_cnt = sigu64_retention_cnt = 0
    sigu64_byte_len = sigu64_retention_byte_len = 0
    sig_invalid_cnt = 0
    sig_get_cnt = sig_getarr_cnt = 0
    sig_get_list = []
    sig_get_list_iso_support = []
    sig_getarr_list = []
    sig_getarr_list_iso_support = []
    sig_invalid_list = []
    sig_enum_get_list = []
    sig_enum_getarr_list = []
    signal_space_max = 0
    signal_enum_space_max = 0
    msg_name_space_max = 0

    for sig in rx_signal_config:
        signal_space_max = max(signal_space_max, len(sig['SignalName']))
        signal_enum_space_max = max(signal_enum_space_max, len(sig['VSignalEnum']))
        msg_name_space_max = max(msg_name_space_max, len(sig['MsgName']))

        bits = sig['Lengthinbits']
        bytes_len = sig['LengthInBytes']

        if bits <= 8:
            if sig['ValueRetention']:
                sigu8_retention_cnt += 1
            else:
                sigu8_cnt += 1
        elif bits <= 16:
            if sig['ValueRetention']:
                sigu16_retention_cnt += 1
            else:
                sigu16_cnt += 1
        elif bits <= 32:
            if sig['ValueRetention']:
                sigu32_retention_cnt += 1
            else:
                sigu32_cnt += 1
        else:
            if sig['ValueRetention']:
                sigu64_retention_cnt += 1
                sigu64_retention_byte_len += bytes_len
            else:
                sigu64_cnt += 1
                sigu64_byte_len += bytes_len

        if sig['SignalInvalidSupport']:
            sig_invalid_cnt += 1
            sig_invalid_list.append(sig['SignalName'])

        if bytes_len <= 4:
            sig_get_cnt += 1
            sig_get_list.append(sig['SignalName'])
            sig_enum_get_list.append(sig['VSignalEnum'])
            sig_get_list_iso_support.append(sig['SignalIsoSupport'])
        else:
            sig_getarr_cnt += 1
            sig_getarr_list.append(sig['SignalName'])
            sig_enum_getarr_list.append(sig['VSignalEnum'])
            sig_getarr_list_iso_support.append(sig['SignalIsoSupport'])

    # ==================== Generate .h file ====================
    with open(HEADER_FILE, "w") as f:
        f.write(f'''/********************************************************************************************************************
* C O P Y R I G H T                                                                                                 *
*********************************************************************************************************************
* Copyright (c) 2019 by Visteon Corporation.       All rights reserved.                                             *
*                                                                                                                   *
* NOTICE: This Is an unpublished work Of authorship, which contains trade secrets.                                  *
* Visteon Corporation owns all rights to this work And intends to maintain it in confidence to preserve             *
* its trade secret status. Visteon Corporation reserves the right, under the copyright laws of the United States    *
* Or those of any other country that may have jurisdiction, to protect this work as an unpublished work,            *
* in the event of an inadvertent Or deliberate unauthorized publication. Visteon Corporation also reserves its      *
* rights under all copyright laws to protect this work as a published work, when appropriate.                       *
* Those having access to this work may Not copy it, use it, modify it, Or disclose the information contained in it  *
* without the written authorization of Visteon Corporation.                                                         *
********************************************************************************************************************/
/*!*****************************************************************************************************************
*    @file {HEADER_FILE}
*    @ingroup {MODULE_NAME}
*    @brief This file contains the configuration of the {MODULE_NAME} component.
*    Auto generated code in {user_name} at {gen_time}
********************************************************************************************************************/
#ifndef {PREFIX}_RX_CFG_H_
#define {PREFIX}_RX_CFG_H_

/********************************************************************************************************************
*  PUBLIC MACRO DEFINITIONS                                                                                         *
********************************************************************************************************************/
#include "Rte_C{MODULE_NAME}_Type.h"
#include "{MODULE_NAME}.h"
#include "Rte_Type.h"

#define {PREFIX}_SEC_CODE_START
#define {PREFIX}_CORE_CONST_SEC_START
#define {PREFIX}_CORE_DATA_SEC_START
#define {PREFIX}_CORE_BSS_SEC_START
#include "MemMap.h"

/********************************************************************************************************************
*  PUBLIC MACRO DEFINITIONS                                                                                         *
********************************************************************************************************************/
''')

        if general_iso_support:
            f.write(f'#define {PREFIX}_RXCAN_ISO_SIGNAL_SUPPORTED\n')

        space_length = 60

        def pad(name):
            return " " * (space_length - len(name))

        # Buffer counts
        f.write('''/* Number of signal buffers */\n''')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U8_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U8_BUFFERS")} {max(sigu8_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U16_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U16_BUFFERS")} {max(sigu16_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U32_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U32_BUFFERS")} {max(sigu32_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U8ARR_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U8ARR_BUFFERS")} {max(sigu64_byte_len, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U8ARR_LENGTH_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U8ARR_LENGTH_BUFFERS")} {max(sigu64_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U8_RETENTION_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U8_RETENTION_BUFFERS")} {max(sigu8_retention_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U16_RETENTION_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U16_RETENTION_BUFFERS")} {max(sigu16_retention_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U32_RETENTION_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U32_RETENTION_BUFFERS")} {max(sigu32_retention_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U8ARR_RETENTION_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_U8ARR_RETENTION_BUFFERS")} {max(sigu64_retention_byte_len, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_RETENTION_U8ARR_LENGTH_BUFFERS{pad(f"{PREFIX}_RXCAN_NUM_OF_RETENTION_U8ARR_LENGTH_BUFFERS")} {max(sigu64_retention_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_U8ARR_SIG_CFG{pad(f"{PREFIX}_RXCAN_NUM_OF_U8ARR_SIG_CFG")} {max(sig_getarr_cnt, 1)}U\n')

        # Status bytes
        nr_bytes = (len(rx_signal_config) + 7) // 8 or 1
        f.write(f'\n/* Number of signal status bytes */\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_NR_STATUS_BYTES{pad(f"{PREFIX}_RXCAN_NUM_OF_NR_STATUS_BYTES")} {nr_bytes}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_MISSING_STATUS_BYTES{pad(f"{PREFIX}_RXCAN_NUM_OF_MISSING_STATUS_BYTES")} {nr_bytes}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_JUST_RCVD_STATUS_BYTES{pad(f"{PREFIX}_RXCAN_NUM_OF_JUST_RCVD_STATUS_BYTES")} {nr_bytes}U\n')
        invalid_bytes = (sig_invalid_cnt + 7) // 8 or 1
        f.write(f'#define COMABS_RXCAN_NUM_OF_SIGNAL_INVALID_STATUS_BYTES{pad("COMABS_RXCAN_NUM_OF_SIGNAL_INVALID_STATUS_BYTES")} {invalid_bytes}U\n')
        f.write(f'\n/* Number of invalid signals */\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_INVALID_SIGNALS{pad(f"{PREFIX}_RXCAN_NUM_OF_INVALID_SIGNALS")} {sig_invalid_cnt}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NO_INVALDID_SIGNAL{pad(f"{PREFIX}_RXCAN_NO_INVALDID_SIGNAL")} {PREFIX}_RXCAN_NUM_OF_INVALID_SIGNALS\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_MISSING_SIGNALS{pad(f"{PREFIX}_RXCAN_NUM_OF_MISSING_SIGNALS")} eEndOfRxSignal\n')
        f.write(f'#define {PREFIX}_RXCAN_MAINFUNCTION_TIME_IN_MS{pad(f"{PREFIX}_RXCAN_MAINFUNCTION_TIME_IN_MS")} 10U\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_GET_SIG_FUNCT{pad(f"{PREFIX}_RXCAN_NUM_OF_GET_SIG_FUNCT")} {max(sig_get_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NO_GET_SIG_FUNCT_IDX{pad(f"{PREFIX}_RXCAN_NO_GET_SIG_FUNCT_IDX")} {PREFIX}_RXCAN_NUM_OF_GET_SIG_FUNCT\n')
        f.write(f'#define {PREFIX}_RXCAN_NUM_OF_GET_SIG_U8ARR_FUNCT{pad(f"{PREFIX}_RXCAN_NUM_OF_GET_SIG_U8ARR_FUNCT")} {max(sig_getarr_cnt, 1)}U\n')
        f.write(f'#define {PREFIX}_RXCAN_NO_GET_SIG_U8ARR_FUNCT_IDX{pad(f"{PREFIX}_RXCAN_NO_GET_SIG_U8ARR_FUNCT_IDX")} {PREFIX}_RXCAN_NUM_OF_GET_SIG_U8ARR_FUNCT\n')

        f.write(f'''
#define {PREFIX}_RXCAN_SIG_NR_ENABLE          0x01U
#define {PREFIX}_RXCAN_SIG_MISSING_ENABLE     0x02U
#define {PREFIX}_RXCAN_SIG_INVALID_ENABLE     0x04U
#define {PREFIX}_RXCAN_SIG_RETENTION_ENABLE   0x08U
#define {PREFIX}_RXCAN_SIG_ISO_SUPPORTED      0x10U

#define {PREFIX}_SET_BIT(datau8,bitpos)     ((datau8) |= (uint8)(1u<<(bitpos))) 
#define {PREFIX}_CLEAR_BIT(datau8,bitpos)   ((datau8) &= (uint8)(~(1U << (bitpos))))  
#define {PREFIX}_IS_BIT_SET(datau8,bitpos)  ((( (datau8) & ((uint8)(1u<<(bitpos)))) != 0U) ? TRUE : FALSE) 

/********************************************************************************************************************
*  PUBLIC TYPE DEFINITIONS                                                                                         *
********************************************************************************************************************/
typedef void (*{MODULE_NAME}_RxCanGetSignalFunctPtrType)(uint8 *pSignalValue);
typedef void (*{MODULE_NAME}_RxCanGetSignalU8ArrFunctPtrType)(uint8 *pSignalValue, uint32 *pSignalLength);
typedef boolean (*{MODULE_NAME}_RxCanSignalInvalidCheckFuncPtrType)(void);
typedef uint8 {MODULE_NAME}_RxCanSigMonitoringType;

typedef struct
{{
    uint16 SignalLengthInByte;
    uint16 SignalBufferStartIdx;
    uint16 SignalInvalidIdx;
    uint16 SignalGetFunctIdx;
    {MODULE_NAME}_RxCanSigMonitoringType RxSigMonitoringType;
    uint32 SignalMissingCounterValue;
}} {MODULE_NAME}_RxCanSignalConfigType;

typedef struct
{{
    {MODULE_NAME}_RxCanGetSignalFunctPtrType pGetSignalFunct;
}} {MODULE_NAME}_RxCanGetSignalFunctConfigType;

typedef struct
{{
    uint16 SignalU8ArrBufferLenIdx;
    {MODULE_NAME}_RxCanGetSignalU8ArrFunctPtrType pGetSignalU8ArrFunct;
}} {MODULE_NAME}_RxCanGetSignalU8ArrConfigType;

typedef struct
{{
    uint32 SignalInvalidSigIdx;
    uint32 SignalInvalidCounterValue;
    {MODULE_NAME}_RxCanSignalInvalidCheckFuncPtrType pSignalInvalidCheckFunc;
}} {MODULE_NAME}_RxCanSignalInvalidConfigType;

/********************************************************************************************************************
*  PUBLIC CONSTANT AND VARIABLE DECLARATION                                                                        *
********************************************************************************************************************/
extern const {MODULE_NAME}_RxCanGetSignalFunctConfigType {MODULE_NAME}_RxCanGetSignalFunctConfig[{PREFIX}_RXCAN_NUM_OF_GET_SIG_FUNCT];
extern const {MODULE_NAME}_RxCanGetSignalU8ArrConfigType {MODULE_NAME}_RxCanGetSignalU8ArrConfig[{PREFIX}_RXCAN_NUM_OF_U8ARR_SIG_CFG];
extern const {MODULE_NAME}_RxCanSignalConfigType {MODULE_NAME}_RxCanSignalConfig[eEndOfRxSignal];
extern const {MODULE_NAME}_RxCanSignalInvalidConfigType {MODULE_NAME}_RxCanSignalInvalidConfig[{PREFIX}_RXCAN_NUM_OF_INVALID_SIGNALS]; 
extern uint8 {MODULE_NAME}_RxCanSignalNRStatus[{PREFIX}_RXCAN_NUM_OF_NR_STATUS_BYTES];
extern uint8 {MODULE_NAME}_RxCanSignalMissingStatus[{PREFIX}_RXCAN_NUM_OF_MISSING_STATUS_BYTES];
extern uint8 {MODULE_NAME}_RxCanJustRcvdStatus[{PREFIX}_RXCAN_NUM_OF_JUST_RCVD_STATUS_BYTES];

/********************************************************************************************************************
*  PUBLIC FUNCTION PROTOTYPE                                                                                        *
********************************************************************************************************************/
#ifdef {PREFIX}_RXCAN_ISO_SIGNAL_SUPPORTED
/*!******************************************************************************************************************
*   @brief   Main function for the {MODULE_NAME}_RxIso . Should be called in every
*   \\{PREFIX}_RXCAN_MAINFUNCTION_TIME_IN_MS. It process Iso signals
********************************************************************************************************************/
extern void {MODULE_NAME}_RxIsoMainFunction(void);
#endif

/*!******************************************************************************************************************
*   On data rx function for the iso signals
********************************************************************************************************************/
''')

        for i, name in enumerate(sig_get_list):
            if sig_get_list_iso_support[i]:
                f.write(f'extern void C{MODULE_NAME}_OnDataRx_{name}(void);\n')
        for i, name in enumerate(sig_getarr_list):
            if sig_getarr_list_iso_support[i]:
                f.write(f'extern void C{MODULE_NAME}_OnDataRx_{name}(void);\n')

        f.write('''
/*!******************************************************************************************************************
*   Rx Get functions for the signals
********************************************************************************************************************/
''')

        for i, name in enumerate(sig_get_list):
            if not sig_get_list_iso_support[i]:
                sig = next(s for s in rx_signal_config if s['SignalName'] == name)
                bits = sig['Lengthinbits']
                ctype = get_ctype_for_bits(bits)
                f.write(f'extern void {MODULE_NAME}_RxCanGet{name}({ctype} *pSignalValue);\n')

        for i, name in enumerate(sig_getarr_list):
            if not sig_getarr_list_iso_support[i]:
                f.write(f'extern void {MODULE_NAME}_RxCanGet{name}(uint8 *pSignalValue, uint32 *pSignalLength);\n')

        f.write('\n/* Signal Invalid check function */\n')
        for name in sig_invalid_list:
            f.write(f'extern boolean {MODULE_NAME}_RxCanInvalid_{name}(void);\n')

        f.write(f'''
/********************************************************************************************************************
*  Below are the set of signals that can be received via C{MODULE_NAME}_RxCan_GetSignalValue API               			    
********************************************************************************************************************/
/* SignalEnum SignalEnumValue SignalNameInDbc MessageName
''')
        for i, sig in enumerate(rx_signal_config):
            if sig['LengthInBytes'] <= 4:
                f.write(f"   {sig['VSignalEnum']}    {i}     {sig['SignalName']}    {sig['MsgName']}\n")
        f.write("*/\n")

        f.write(f'''
/********************************************************************************************************************
*  Below are the set of signals that can be received via C{MODULE_NAME}_RxCan_GetSignalValueU8Arr API          			    
********************************************************************************************************************/
/* SignalEnum SignalEnumValue SignalNameInDbc MessageName
''')
        for i, sig in enumerate(rx_signal_config):
            if sig['LengthInBytes'] > 4:
                f.write(f"   {sig['VSignalEnum']}    {i}     {sig['SignalName']}    {sig['MsgName']}\n")
        f.write("*/\n")

        f.write(f'''
#define {PREFIX}_SEC_CODE_STOP
#define {PREFIX}_CORE_CONST_SEC_END
#define {PREFIX}_CORE_DATA_SEC_END
#define {PREFIX}_CORE_BSS_SEC_END
#endif /* {PREFIX}_RX_CFG_H_ */
/*******************************************************************************************************************
*                                 END OF FILE                                                                      *
********************************************************************************************************************/
''')

    # ==================== Generate .c file ====================
    with open(SOURCE_FILE, "w") as f:
        f.write(f'''/********************************************************************************************************************
* C O P Y R I G H T                                                                                                 *
*********************************************************************************************************************
* Copyright (c) 2019 by Visteon Corporation.       All rights reserved.                                             *
*                                                                                                                   *
* NOTICE: This Is an unpublished work Of authorship, which contains trade secrets.                                  *
* Visteon Corporation owns all rights to this work And intends to maintain it in confidence to preserve             *
* its trade secret status. Visteon Corporation reserves the right, under the copyright laws of the United States    *
* Or those of any other country that may have jurisdiction, to protect this work as an unpublished work,            *
* in the event of an inadvertent Or deliberate unauthorized publication. Visteon Corporation also reserves its      *
* rights under all copyright laws to protect this work as a published work, when appropriate.                       *
* Those having access to this work may Not copy it, use it, modify it, Or disclose the information contained in it  *
* without the written authorization of Visteon Corporation.                                                         *
********************************************************************************************************************/
/*!*****************************************************************************************************************
*    @file {SOURCE_FILE}
*    @ingroup {MODULE_NAME}
*    @brief This file contains the configuration of the {MODULE_NAME} component.
*    Auto generated code in {user_name} at {gen_time}
********************************************************************************************************************/
#include "{MODULE_NAME}.h"
#include "{HEADER_FILE}"
#include "Rte_C{MODULE_NAME}_Type.h"

#define {PREFIX}_SEC_CODE_START
#define {PREFIX}_CORE_CONST_SEC_START
#define {PREFIX}_CORE_DATA_SEC_START
#define {PREFIX}_CORE_BSS_SEC_START
#include "MemMap.h"

// Global status arrays
uint8 {MODULE_NAME}_RxCanSignalNRStatus[{PREFIX}_RXCAN_NUM_OF_NR_STATUS_BYTES];
uint8 {MODULE_NAME}_RxCanSignal_currentStatus[{PREFIX}_RXCAN_NUM_OF_NR_STATUS_BYTES];
uint8 {MODULE_NAME}_RxCanJustRcvdStatus[{PREFIX}_RXCAN_NUM_OF_JUST_RCVD_STATUS_BYTES];
uint8 {MODULE_NAME}_RxCanSignalMissingStatus[{PREFIX}_RXCAN_NUM_OF_MISSING_STATUS_BYTES];

/********************************************************************************************************************
*  PUBLIC CONSTANT AND VARIABLE DEFINITION                                                                          *
********************************************************************************************************************/
const {MODULE_NAME}_RxCanGetSignalFunctConfigType {MODULE_NAME}_RxCanGetSignalFunctConfig[{PREFIX}_RXCAN_NUM_OF_GET_SIG_FUNCT] =
{{
    /* {{ {MODULE_NAME}_RxCanGetSignalFunctPtrType }} */
''')
        for i, name in enumerate(sig_get_list):
            ptr = "NULL_PTR" if sig_get_list_iso_support[i] else f"&{MODULE_NAME}_RxCanGet{name}"
            comma = "" if i == len(sig_get_list) - 1 else ","
            f.write(f"    {{ {ptr} }}{comma}  /* {i} {sig_enum_get_list[i]} */\n")
        f.write("};\n\n")

        f.write(f'''const {MODULE_NAME}_RxCanGetSignalU8ArrConfigType {MODULE_NAME}_RxCanGetSignalU8ArrConfig[{PREFIX}_RXCAN_NUM_OF_U8ARR_SIG_CFG] =
{{
    /* {{ SignalU8ArrBufferLenIdx, {MODULE_NAME}_RxCanGetSignalU8ArrFunctPtrType }} */
''')
        len_idx = len_ret_idx = 0
        for i, sig in enumerate(rx_signal_config):
            if sig['LengthInBytes'] > 4:
                idx = len_ret_idx if sig['ValueRetention'] else len_idx
                ptr = "NULL_PTR" if sig['SignalIsoSupport'] else f"&{MODULE_NAME}_RxCanGet{sig['SignalName']}"
                comma = "" if i == sig_getarr_cnt - 1 else ","
                f.write(f"    {{ {idx}U, {ptr} }}{comma}  /* {i} {sig['VSignalEnum']} */\n")
                if sig['ValueRetention']:
                    len_ret_idx += 1
                else:
                    len_idx += 1
        f.write("};\n\n")

        # Signal config
        f.write(f'''const {MODULE_NAME}_RxCanSignalConfigType {MODULE_NAME}_RxCanSignalConfig[eEndOfRxSignal] =
{{
    /* {{ SignalLengthInByte , SignalBufferStartIdx , SignalInvalidIdx , SignalGetFunctIdx , */
    /*   RxSigMonitoringType , SignalMissingCounterValue }} */
''')
        l_u8 = l_u8_ret = l_u16 = l_u16_ret = l_u32 = l_u32_ret = l_u64 = l_u64_ret = 0
        get_id = 0
        get_id_arr = 0
        inv_id = 0
        for i, sig in enumerate(rx_signal_config):
            bits = sig['Lengthinbits']
            bytes_len = sig['LengthInBytes']
            if bits <= 8:
                buf_idx = l_u8 if not sig['ValueRetention'] else l_u8_ret
                if sig['ValueRetention']: l_u8_ret += 1
                else: l_u8 += 1
            elif bits <= 16:
                buf_idx = l_u16 if not sig['ValueRetention'] else l_u16_ret
                if sig['ValueRetention']: l_u16_ret += 1
                else: l_u16 += 1
            elif bits <= 32:
                buf_idx = l_u32 if not sig['ValueRetention'] else l_u32_ret
                if sig['ValueRetention']: l_u32_ret += 1
                else: l_u32 += 1
            else:
                buf_idx = l_u64 if not sig['ValueRetention'] else l_u64_ret
                if sig['ValueRetention']: l_u64_ret += bytes_len
                else: l_u64 += bytes_len

            inv_idx = str(inv_id) if sig['SignalInvalidSupport'] else f"{PREFIX}_RXCAN_NO_INVALDID_SIGNAL"
            if sig['SignalInvalidSupport']: inv_id += 1

            get_idx = get_id if bytes_len <= 4 else get_id_arr
            if bytes_len <= 4: get_id += 1
            else: get_id_arr += 1

            flags = [f"{PREFIX}_RXCAN_SIG_NR_ENABLE"]
            if sig['TimeoutValue'] > 0:
                flags.append(f"{PREFIX}_RXCAN_SIG_MISSING_ENABLE")
            if sig['SignalInvalidSupport']:
                flags.append(f"{PREFIX}_RXCAN_SIG_INVALID_ENABLE")
            if sig['ValueRetention']:
                flags.append(f"{PREFIX}_RXCAN_SIG_RETENTION_ENABLE")
            if sig['SignalIsoSupport']:
                flags.append(f"{PREFIX}_RXCAN_SIG_ISO_SUPPORTED")
            flag_str = " | ".join(flags)

            comma = "" if i == len(rx_signal_config) - 1 else ","
            f.write(f"    {{{bytes_len}U, {buf_idx}U, {inv_idx}, {get_idx}U, ({flag_str}), {int(sig['TimeoutValue']/10)}U}}{comma}  /* {i} {sig['SignalName']} {sig['VSignalEnum']} */\n")
        f.write("};\n\n")

        # Invalid config
        f.write(f'''const {MODULE_NAME}_RxCanSignalInvalidConfigType {MODULE_NAME}_RxCanSignalInvalidConfig[{PREFIX}_RXCAN_NUM_OF_INVALID_SIGNALS] =
{{
    /* {{ SignalInvalidSigIdx, SignalInvalidCounterValue, pSignalInvalidCheckFunc }} */
''')
        if sig_invalid_cnt > 0:
            for i, sig in enumerate(rx_signal_config):
                if sig['SignalInvalidSupport']:
                    counter = int(sig['SignalInvalidCounter'] / 10)
                    func = f"&{MODULE_NAME}_RxCanInvalid_{sig['SignalName']}"
                    comma = "" if i == sig_invalid_cnt - 1 else ","
                    f.write(f"    {{ {sig['VSignalEnum']}, {counter}U, {func} }}{comma}  /* {i} {sig['SignalName']} */\n")
        else:
            f.write("    { 0U, 0U, NULL_PTR }\n")
        f.write("};\n\n")

        # Getter functions
        for sig in rx_signal_config:
            name = sig['SignalName']
            bits = sig['Lengthinbits']
            bytes_len = sig['LengthInBytes']
            if bytes_len <= 4:
                ctype = get_ctype_for_bits(bits)
                f.write(f'''
/********************************************************************************************************************
*   @brief   This method is used to get the signal {name}
*   @param[in] pSignalValue pointer to signal data
********************************************************************************************************************/
void {MODULE_NAME}_RxCanGet{name}({ctype} *pSignalValue)
{{
    (void)Rte_Read_rp_SR_Rx{name}_{name}(pSignalValue);
}}
''')
            else:
                f.write(f'''
/********************************************************************************************************************
*   @brief   This method is used to get the signal {name}
*   @param[in] pSignalValue pointer to signal data
*   @param[in] pSignalLength length of the data
********************************************************************************************************************/
void {MODULE_NAME}_RxCanGet{name}(uint8 *pSignalValue, uint32 *pSignalLength)
{{
''')
                if sig['DynSignalSupport']:
                    f.write(f'    uint16 SigLength = 0U;\n')
                    f.write(f'    (void)Rte_Receive_rp_SR_Rx{name}_{name}(pSignalValue, &SigLength);\n')
                    f.write(f'    *pSignalLength = SigLength;\n')
                else:
                    f.write(f'    (void)Rte_Read_rp_SR_Rx{name}_{name}(pSignalValue);\n')
                    f.write(f'    *pSignalLength = {bytes_len}U;\n')
                f.write("}\n")

        # OnDataRx for UB and message callbacks
        rx_msg_sig_list = {}
        rx_msg_list = []
        rx_msg_key_sig_map = {}
        rx_Ub_sig_list = []
        for sig in rx_signal_config:
            msg_name = sig['MsgName']
            if sig['SignalUBSupport']:
                rx_Ub_sig_list.append(sig)
            else:
                if msg_name not in rx_msg_list:
                    rx_msg_list.append(msg_name)
                    if sig['MessageCbk']:
                        rx_msg_key_sig_map[msg_name] = sig['SignalName']
                    rx_msg_sig_list[msg_name] = []
                rx_msg_sig_list[msg_name].append(sig)

        for sig in rx_Ub_sig_list:
            f.write(f'''
/********************************************************************************************************************
*   @brief   This method is called from COM on successful reception of the signal {sig['SignalName']}
********************************************************************************************************************/
void C{MODULE_NAME}_OnDataRx_{sig['SignalName']}(void)
{{
    {PREFIX}_CLEAR_BIT({MODULE_NAME}_RxCanSignalNRStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));
    {PREFIX}_SET_BIT({MODULE_NAME}_RxCanJustRcvdStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));
    {PREFIX}_SET_BIT({MODULE_NAME}_RxCanSignal_currentStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));
    {MODULE_NAME}_ReloadMissingTimer({sig['VSignalEnum']});
}}
''')

        for msg in rx_msg_list:
            if msg in rx_msg_key_sig_map:
                f.write(f'''
/********************************************************************************************************************
*   @brief   This method is called from COM on successful reception of the message {msg}
********************************************************************************************************************/
void C{MODULE_NAME}_OnDataRx_{rx_msg_key_sig_map[msg]}(void)
{{
''')
                for sig in rx_msg_sig_list[msg]:
                    f.write(f"    {PREFIX}_CLEAR_BIT({MODULE_NAME}_RxCanSignalNRStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));\n")
                    f.write(f"    {PREFIX}_SET_BIT({MODULE_NAME}_RxCanJustRcvdStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));\n")
                    f.write(f"    {PREFIX}_SET_BIT({MODULE_NAME}_RxCanSignal_currentStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));\n")
                    f.write(f"    {MODULE_NAME}_ReloadMissingTimer({sig['VSignalEnum']});\n")
                f.write("}\n")

        # Invalid check functions
        for sig in rx_signal_config:
            if sig['LengthInBytes'] <= 4 and sig['SignalInvalidSupport']:
                ctype = get_ctype_for_bits(sig['Lengthinbits'])
                f.write(f'''
/********************************************************************************************************************
*   @brief   This method is used to validate the signal {sig['SignalName']}
********************************************************************************************************************/
boolean {MODULE_NAME}_RxCanInvalid_{sig['SignalName']}(void)
{{
    {ctype} SignalValue = 0U;
    boolean Status = FALSE;
''')
                qf_sig = rx_signal_qf_invalid_list.get(sig['SignalQfSignalValue']) if sig['SignalQfSupport'] else None
                src_sig = qf_sig if qf_sig else sig
                f.write(f'    {MODULE_NAME}_RxCanGet{src_sig["SignalName"]}(&SignalValue);\n')
                f.write(f'    /* Invalid value: {src_sig["SignalInvalidVal"]} */\n')
                vals = str(src_sig['SignalInvalidVal']).split('|')
                cond = " || ".join([
                    f"(SignalValue >= {v.strip('()').split('-')[0]}U && SignalValue <= {v.strip('()').split('-')[1]}U)" if '-' in v
                    else f"(SignalValue == {v.strip('()')}U)"
                    for v in vals
                ])
                f.write(f'    if ({cond})\n    {{\n        Status = TRUE;\n    }}\n')
                f.write('    return Status;\n}\n')

        # ISO callbacks
        for sig in rx_signal_config:
            if sig['SignalIsoSupport']:
                f.write(f'''
/********************************************************************************************************************
*   @brief   This method is called from COM on successful reception of the ISO signal {sig['SignalName']}
********************************************************************************************************************/
void C{MODULE_NAME}_OnDataRx_{sig['SignalName']}(void)
{{
    {PREFIX}_CLEAR_BIT({MODULE_NAME}_RxCanSignalNRStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));
    {PREFIX}_SET_BIT({MODULE_NAME}_RxCanJustRcvdStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));
    {PREFIX}_SET_BIT({MODULE_NAME}_RxCanSignal_currentStatus[{sig['VSignalEnum']} >> 3U], ({sig['VSignalEnum']} & 0x7U));
    {MODULE_NAME}_ReloadMissingTimer({sig['VSignalEnum']});
}}
''')

        f.write('''
#define {PREFIX}_SEC_CODE_STOP
#define {PREFIX}_CORE_CONST_SEC_END
#define {PREFIX}_CORE_DATA_SEC_END
#define {PREFIX}_CORE_BSS_SEC_END
#include "MemMap.h"
/*******************************************************************************************************************
*                                 END OF FILE                                                                      *
********************************************************************************************************************/
''')

# ================ MAIN ================
if __name__ == '__main__':
    excel_name = "E4.0_Intermediate_sheet.xlsx"
    if len(sys.argv) > 1:
        excel_name = sys.argv[1]
    current_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    username = getpass.getuser()
    generate_code_comabs_rx_cfg(excel_name, username, current_time)