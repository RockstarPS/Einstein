# main.py
import subprocess
import click
import ComAbsMdlCodeGen
import ComAbsMdlTxCodeGen
import ComAbsMdlRxCodeGen
import datetime
import getpass
import os
from openpyxl import load_workbook

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_DBC_DIR = os.path.join(SCRIPT_DIR, "..", "Input_DBC")
INTERMEDIATE_SHEET_DIR = os.path.join(SCRIPT_DIR, "..", "Intermediate_Sheet")
GEN_FILES_DIR = os.path.join(SCRIPT_DIR, "..", "gen-files")
CONTEXT_SETTINGS = dict(help_option_names=['-h', '--help'])

def validate_excel_sheets(excel_path: str):
    """Validate required fields in both RxSignals and TxSignals sheets."""
    try:
        wb = load_workbook(excel_path, data_only=True)
    except Exception as e:
        raise FileNotFoundError(f"Failed to open Excel file: {e}")

    errors = []
    warnings = []

    # Validate RxSignals sheet
    try:
        rx_sheet = wb["RxSignals"]
        sheet_errors, sheet_warnings = validate_sheet(rx_sheet, "RxSignals")
        errors.extend(sheet_errors)
        warnings.extend(sheet_warnings)
    except KeyError:
        errors.append("Sheet 'RxSignals' not found in Excel file")

    # Validate TxSignals sheet
    try:
        tx_sheet = wb["TxSignals"]
        sheet_errors, sheet_warnings = validate_sheet(tx_sheet, "TxSignals")
        errors.extend(sheet_errors)
        warnings.extend(sheet_warnings)
    except KeyError:
        errors.append("Sheet 'TxSignals' not found in Excel file")

    # Print warnings (e.g., disabled or blank Enable/Disable)
    for warn in warnings:
        print(warn)

    # Only stop if there are actual validation errors (excluding Enable/Disable)
    if errors:
        print("\n🛑 VALIDATION FAILED: The script cannot proceed due to the following issues:")
        for err in errors:
            print(f"  - {err}")
        raise ValueError("Excel validation failed — fix the errors above.")

    print("✅ Excel validation passed. Proceeding with code generation...")

def validate_sheet(sheet, sheet_name: str) -> tuple[list[str], list[str]]:
    """Validate required fields and return (errors, warnings) for disabled signals."""
    errors = []
    warnings = []

    # Required columns (excluding 'Enable/Disable' from strict validation)
    required_cols = {
        "Signal Name": 1,
        "VSignal Name": 2,
        "VSignal Enum": 3,
        "Length in bits": 5,
        "Message Name": 6 if sheet_name == "TxSignals" else 7,
    }

    for row in range(2, sheet.max_row + 1):
        signal_cell = sheet.cell(row=row, column=1)
        signal_name = signal_cell.value or f"Unknown_Signal_Row{row}"

        msg_cell = sheet.cell(row=row, column=required_cols["Message Name"])
        msg_name = msg_cell.value or f"Unknown_Msg_Row{row}"

        # Check only the truly required fields
        for col_name, col_idx in required_cols.items():
            cell = sheet.cell(row=row, column=col_idx)
            if not cell.value or str(cell.value).strip() == "":
                errors.append(f"Row {row}, {sheet_name}, Column '{col_name}': Value is missing or blank")

        # Skip further checks if critical fields are missing
        if any(f"Row {row}" in err for err in errors):
            continue

        # Validate Length in bits
        length_cell = sheet.cell(row=row, column=5)
        length_val = length_cell.value
        try:
            length_num = int(length_val)
            if length_num <= 0:
                errors.append(f"Row {row}, {sheet_name}, Column 'Length in bits': Must be greater than 0")
        except (ValueError, TypeError):
            errors.append(f"Row {row}, {sheet_name}, Column 'Length in bits': Invalid value '{length_val}'")

        # Special validation for RxSignals
        if sheet_name == "RxSignals":
            timeout_cell = sheet.cell(row=row, column=6)
            timeout_val = timeout_cell.value
            if timeout_val is None or (isinstance(timeout_val, str) and timeout_val.strip() == ""):
                errors.append(f"Row {row}, {sheet_name}, Column 'Timeout Value in ms': Value is not configured")
            else:
                try:
                    timeout_num = int(timeout_val)
                    if timeout_num <= 0:
                        errors.append(f"Row {row}, {sheet_name}, Column 'Timeout Value in ms': Must be greater than 0")
                except (ValueError, TypeError):
                    errors.append(f"Row {row}, {sheet_name}, Column 'Timeout Value in ms': Invalid value '{timeout_val}'")

        # 🔍 Check for disabled signals (Enable/Disable == FALSE)
        enable_cell = sheet.cell(row=row, column=4)
        enable_value = str(enable_cell.value).strip().upper() if enable_cell.value is not None else ""
        if enable_value == "FALSE":
            warnings.append(f"The signal is Disabled: {signal_name} (Msg: {msg_name})")
        elif enable_value == "":
            warnings.append(f"Warning: Enable/Disable is blank for signal '{signal_name}' (Msg: {msg_name})")

    return errors, warnings

@click.group(context_settings=CONTEXT_SETTINGS)
def main():
    pass

@main.command()
@click.option('--iexcel', '-iexcel', type=str, help="Input SRD Excel", required=True)
@click.option('--oexcel', '-oexcel', type=str, help="Output Excel", required=True)
@click.option('--dbc', '-dbc', type=str, help="network:file,network:file", required=True)
@click.option('--node', '-node', type=str, help="ECU node", required=True)
@click.option('--inifile', '-inifile', type=str, help="Config INI", required=True)
def genexcel(iexcel, oexcel, dbc, node, inifile):
    input_srd_path = iexcel if os.path.isabs(iexcel) else os.path.join(SCRIPT_DIR, iexcel)
    if not os.path.isfile(input_srd_path):
        raise FileNotFoundError(f"Input SRD Excel not found: {input_srd_path}")

    dbc_list = []
    for item in dbc.split(','):
        try:
            network_name, dbc_filename = item.split(':')
        except ValueError:
            raise ValueError(f"Invalid DBC format: {item}. Use 'network:filename'")
        dbc_path = os.path.join(INPUT_DBC_DIR, dbc_filename)
        if not os.path.isfile(dbc_path):
            raise FileNotFoundError(f"DBC not found: {dbc_path}")
        dbc_list.append({'network_name': network_name, 'dbc_name': dbc_path})

    output_excel_path = os.path.join(INTERMEDIATE_SHEET_DIR, oexcel)
    os.makedirs(INTERMEDIATE_SHEET_DIR, exist_ok=True)

    print("ComAbsMdl configuration Excel generation started")
    ComAbsMdlExcelGen.update_ini_cfg(inifile)
    ComAbsMdlExcelGen.generate_comabs_cfg_excel_sheet(input_srd_path, dbc_list[0]['dbc_name'], node, output_excel_path)
    print(f"ComAbsMdl configuration Excel generated: {output_excel_path}")

@main.command()
@click.option('--iexcel', '-iexcel', type=str, help="Input Excel file", required=True)
@click.option('--safe-rx', is_flag=True, help="Enable Safe Rx")
def gencode(iexcel, safe_rx):
    current_time = datetime.datetime.now()
    username = getpass.getuser()
    output_dir = "..\\gen-files"
    input_excel_path = os.path.join(INTERMEDIATE_SHEET_DIR, iexcel)
    os.makedirs(output_dir, exist_ok=True)

    if not os.path.isfile(input_excel_path):
        raise FileNotFoundError(f"Input Excel file not found: {input_excel_path}")

    print("Validating Excel sheets...")
    validate_excel_sheets(input_excel_path)

    print("ComAbsMdl configuration code generation started")
    ComAbsMdlCodeGen.generate_code_comabsmdl(username, current_time, output_dir=".")
    ComAbsMdlTxCodeGen.generate_code_comabs_tx_cfg(input_excel_path, username, current_time, output_dir=".")
    ComAbsMdlRxCodeGen.generate_code_comabs_rx_cfg(input_excel_path, username, current_time, output_dir=".")

    if safe_rx:
        print("Safe Rx function generation enabled")

    print("ComAbsMdl configuration code generation completed")

if __name__ == '__main__':
    uint_testing = False
    try:
        subprocess.call(["taskkill", "/f", "/im", "EXCEL.EXE"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    finally:
        pass

    if uint_testing:
        from timeit import default_timer as timer
        comabs_excel = "E4.0_Intermediate_sheet.xlsx"
        input_dbc = "PlatformDBCE4v7.dbc"
        comabs_excel_path = os.path.join(INTERMEDIATE_SHEET_DIR, comabs_excel)
        dbc_path = os.path.join(INPUT_DBC_DIR, input_dbc)
        node = "IC"
        current_time = datetime.datetime.now()
        username = getpass.getuser()
        output_dir = "..\\gen-files"
        os.makedirs(output_dir, exist_ok=True)

        if not os.path.isfile(comabs_excel_path):
            raise FileNotFoundError(f"Test Excel not found: {comabs_excel_path}")
        if not os.path.isfile(dbc_path):
            raise FileNotFoundError(f"Test DBC not found: {dbc_path}")

        start = timer()
        ComAbsMdlCodeGen.generate_code_comabsmdl(username, current_time, output_dir=".")
        ComAbsMdlTxCodeGen.generate_code_comabs_tx_cfg(comabs_excel_path, username, current_time, output_dir=".")
        ComAbsMdlRxCodeGen.generate_code_comabs_rx_cfg(comabs_excel_path, username, current_time, output_dir=".")
        end = timer()
        print("Generated in ", str(end - start), " seconds")
    else:
        main()