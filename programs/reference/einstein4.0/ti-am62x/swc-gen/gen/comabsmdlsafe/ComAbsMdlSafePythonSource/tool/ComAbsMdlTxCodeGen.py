# ComAbsMdlTxCodeGen.py
import openpyxl
import sys
import datetime
import getpass
import os
from typing import Dict, List, Any

# ==================== CONFIGURATION ====================
MODULE_NAME = "ComAbsMdlSafe"
HEADER_FILE = f"{MODULE_NAME}_TxCfg.h"
SOURCE_FILE = f"{MODULE_NAME}_TxCfg.c"
PREFIX = MODULE_NAME.upper()
ENUM_END = "eEndofSafeTxSignals"  # Must match enum in Excel
def get_string_from_bytes(val) -> str:
    if isinstance(val, bytes):
        return val.decode('utf-8').strip()
    elif val is None:
        return ""
    else:
        return str(val).strip()

def get_bool_val(val) -> bool:
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    return str(val).strip().upper() == "TRUE"
def load_tx_signals(excel_path: str, sheet_name: str = "TxSignals") -> List[Dict[str, Any]]:
    """Load enabled Tx signals from Excel."""
    try:
        workbook = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        print(f"Error opening Excel file: {e}")
        sys.exit(1)

    try:
        sheet = workbook[sheet_name]
    except KeyError:
        print(f"Error: Sheet '{sheet_name}' not found in workbook.")
        sys.exit(1)

    signals = []
    seen_enums = set()

    for row in range(2, sheet.max_row + 1):
        if not get_bool_val(sheet.cell(row=row, column=4).value):  # Enabled
            continue

        try:
            enum_val = get_string_from_bytes(sheet.cell(row=row, column=3).value)
            if not enum_val:
                print(f"Warning: Empty VSignalEnum at row {row}, skipping.")
                continue
            if enum_val in seen_enums:
                raise ValueError(f"Duplicate VSignalEnum '{enum_val}' at row {row}")
            seen_enums.add(enum_val)

            signal = {
                'SignalName': get_string_from_bytes(sheet.cell(row=row, column=1).value),
                'VSignalName': get_string_from_bytes(sheet.cell(row=row, column=2).value),
                'VSignalEnum': enum_val,
                'Lengthinbits': int(sheet.cell(row=row, column=5).value or 0),
                'MsgName': get_string_from_bytes(sheet.cell(row=row, column=6).value),
                'InternalSignal': get_bool_val(sheet.cell(row=row, column=7).value),
                'TxConfirm': get_bool_val(sheet.cell(row=row, column=8).value),
                'ValueRetention': get_bool_val(sheet.cell(row=row, column=9).value),
                'SignalIsoSupport': get_bool_val(sheet.cell(row=row, column=10).value),
            }
            signal['LengthInBytes'] = (signal['Lengthinbits'] + 7) // 8
            signals.append(signal)
        except Exception as e:
            print(f"Error parsing row {row}: {e}")
            sys.exit(1)
    return signals
def load_tx_signals(excel_path: str, sheet_name: str = "TxSignals") -> List[Dict[str, Any]]:
    """Load enabled Tx signals from Excel."""
    try:
        workbook = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        print(f"Error opening Excel file: {e}")
        sys.exit(1)

    try:
        sheet = workbook[sheet_name]
    except KeyError:
        print(f"Error: Sheet '{sheet_name}' not found in workbook.")
        sys.exit(1)

    signals = []
    seen_enums = set()

    for row in range(2, sheet.max_row + 1):
        if not get_bool_val(sheet.cell(row=row, column=4).value):  # Enabled
            continue

        try:
            enum_val = get_string_from_bytes(sheet.cell(row=row, column=3).value)
            if not enum_val:
                print(f"Warning: Empty VSignalEnum at row {row}, skipping.")
                continue
            if enum_val in seen_enums:
                raise ValueError(f"Duplicate VSignalEnum '{enum_val}' at row {row}")
            seen_enums.add(enum_val)

            signal = {
                'SignalName': get_string_from_bytes(sheet.cell(row=row, column=1).value),
                'VSignalName': get_string_from_bytes(sheet.cell(row=row, column=2).value),
                'VSignalEnum': enum_val,
                'Lengthinbits': int(sheet.cell(row=row, column=5).value or 0),
                'MsgName': get_string_from_bytes(sheet.cell(row=row, column=6).value),
                'InternalSignal': get_bool_val(sheet.cell(row=row, column=7).value),
                'TxConfirm': get_bool_val(sheet.cell(row=row, column=8).value),
                'ValueRetention': get_bool_val(sheet.cell(row=row, column=9).value),
                'SignalIsoSupport': get_bool_val(sheet.cell(row=row, column=10).value),
            }
            signal['LengthInBytes'] = (signal['Lengthinbits'] + 7) // 8
            signals.append(signal)
        except Exception as e:
            print(f"Error parsing row {row}: {e}")
            sys.exit(1)
    return signals
def analyze_tx_signals(signals: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Analyze signals for buffer counts and function lists."""
    sigu8_cnt = sigu16_cnt = sigu32_cnt = siguarr_cnt = 0
    siguarr_byte_len = 0
    set_sig = set_sigArr = 0
    confirm_sig_list = []
    confirm_iso_sig_list = []
    set_sig_list = []
    set_sig_list_u8Arr = []
    set_sig_list_props = {}
    set_sig_list_u8Arr_props = {}
    confirm_msg_list = {}

    for sig in signals:
        msg = sig['MsgName']
        if sig['TxConfirm']:
            if msg not in confirm_msg_list:
                confirm_msg_list[msg] = []
            confirm_msg_list[msg].append(sig['SignalName'])
            confirm_sig_list.append(sig['SignalName'])
            confirm_iso_sig_list.append(sig['SignalIsoSupport'])

        if sig['LengthInBytes'] <= 4:
            set_sig += 1
            set_sig_list.append(sig['SignalName'])
            set_sig_list_props[sig['SignalName']] = sig['VSignalEnum']
        else:
            set_sigArr += 1
            set_sig_list_u8Arr.append(sig['SignalName'])
            set_sig_list_u8Arr_props[sig['SignalName']] = sig['VSignalEnum']

        if sig['ValueRetention']:
            if sig['Lengthinbits'] <= 8:
                sigu8_cnt += 1
            elif sig['Lengthinbits'] <= 16:
                sigu16_cnt += 1
            elif sig['Lengthinbits'] <= 32:
                sigu32_cnt += 1
            else:
                siguarr_cnt += 1
                siguarr_byte_len += sig['LengthInBytes']

    return {
        'sigu8_cnt': max(sigu8_cnt, 1),
        'sigu16_cnt': max(sigu16_cnt, 1),
        'sigu32_cnt': max(sigu32_cnt, 1),
        'siguarr_byte_len': max(siguarr_byte_len, 1),
        'siguarr_cnt': max(siguarr_cnt, 1),
        'set_sig': set_sig,
        'set_sigArr': set_sigArr,
        'confirm_sig_list': confirm_sig_list,
        'confirm_iso_sig_list': confirm_iso_sig_list,
        'set_sig_list': set_sig_list,
        'set_sig_list_u8Arr': set_sig_list_u8Arr,
        'set_sig_list_props': set_sig_list_props,
        'set_sig_list_u8Arr_props': set_sig_list_u8Arr_props,
        'confirm_msg_list': confirm_msg_list,
        'total_signals': len(signals),
        'general_iso_support': any(s['SignalIsoSupport'] for s in signals),
    }

def generate_tx_header_file(signals: List[Dict[str, Any]], analysis: Dict[str, Any], user_name: str, gen_time: str, output_dir: str):
    header_path = os.path.join("..\gen-files", HEADER_FILE)
    os.makedirs(output_dir, exist_ok=True)
    with open(header_path, "w") as f:
        f.write('''/********************************************************************************************************************
 * C O P Y R I G H T
 ********************************************************************************************************************
 * Copyright (c) 2019 by Visteon Corporation.       All rights reserved.
 *
 * NOTICE: This Is an unpublished work Of authorship, which contains trade secrets.
 * Visteon Corporation owns all rights to this work And intends to maintain it in confidence to preserve
 * its trade secret status. Visteon Corporation reserves the right, under the copyright laws of the United States
 * Or those of any other country that may have jurisdiction, to protect this work as an unpublished work,
 * in the event of an inadvertent Or deliberate unauthorized publication. Visteon Corporation also reserves its
 * rights under all copyright laws to protect this work as a published work, when appropriate.
 * Those having access to this work may Not copy it, use it, modify it, Or disclose the information contained in it
 * without the written authorization of Visteon Corporation.
 ********************************************************************************************************************/
''')

        f.write(f'''/*!*****************************************************************************************************************
 *    @file {HEADER_FILE}
 *    @ingroup {MODULE_NAME}
 *    @brief This file contains the configuration of the {MODULE_NAME} component.
 *    Auto generated code in {user_name} at {gen_time}
 ********************************************************************************************************************/
#ifndef {PREFIX}_TX_CFG_H_
#define {PREFIX}_TX_CFG_H_

#include "Rte_CComAbsMdlSafe_Type.h"
#include "ComAbsMdlSafe.h"
#include "Rte_Type.h"

#define {PREFIX}_SEC_CODE_START
#define {PREFIX}_CORE_CONST_SEC_START
#define {PREFIX}_CORE_DATA_SEC_START
#define {PREFIX}_CORE_BSS_SEC_START
#include "MemMap.h"

/********************************************************************************************************************
 *  PUBLIC MACRO DEFINITIONS
 ********************************************************************************************************************/
''')

        if analysis['general_iso_support']:
            f.write(f"#define {PREFIX}_TXCAN_ISO_SIGNAL_SUPPORTED\n")

        f.write(f"#define {PREFIX}_TXCAN_NUM_OF_U8_BUFFERS            {analysis['sigu8_cnt']}U\n")
        f.write(f"#define {PREFIX}_TXCAN_NUM_OF_U16_BUFFERS           {analysis['sigu16_cnt']}U\n")
        f.write(f"#define {PREFIX}_TXCAN_NUM_OF_U32_BUFFERS           {analysis['sigu32_cnt']}U\n")
        f.write(f"#define {PREFIX}_TXCAN_NUM_OF_U8ARR_BUFFERS         {analysis['siguarr_byte_len']}U\n")
        f.write(f"#define {PREFIX}_TXCAN_NUM_OF_U8ARR_SIGNAL          {analysis['siguarr_cnt']}U\n")
        f.write(f"#define {PREFIX}_TXCAN_NUM_OF_CONF_STATUS_BYTES     {((analysis['total_signals'] + 7) // 8) or 1}U\n")
        f.write(f"#define CAN{PREFIX}_NUM_OF_SET_SIG_FUNCT            {max(analysis['set_sig'], 1)}U\n")
        f.write(f"#define {PREFIX}_TXCAN_NO_SET_SIG_FUNCT              CAN{PREFIX}_NUM_OF_SET_SIG_FUNCT\n")
        f.write(f"#define CAN{PREFIX}_NUM_OF_SET_SIG_U8ARR_FUNCT       {max(analysis['set_sigArr'], 1)}U\n")
        f.write(f"#define {PREFIX}_TXCAN_NO_SET_SIG_U8ARR_FUNCT         CAN{PREFIX}_NUM_OF_SET_SIG_U8ARR_FUNCT\n")

        f.write(f'''
#define {PREFIX}_TXCAN_SIG_STATUS_NONE               0x0U
#define {PREFIX}_TXCAN_KEEP_ALIVE_BIT_ENABLED        0x1U
#define {PREFIX}_TXCAN_CONFIRMATION_ENABLED          0x2U
#define {PREFIX}_TXCAN_ISO_SUPPORTED                 0x4U
''')

        f.write(f'''
/********************************************************************************************************************
 *  PUBLIC TYPE DEFINITIONS
 ********************************************************************************************************************/
typedef Std_ReturnType (*{MODULE_NAME}_TxCanSetSignalFunctPtrType)(uint8 *pSignalValue);
typedef Std_ReturnType (*{MODULE_NAME}_TxCanSetSignalU8ArrFunctPtrType)(uint8 *pSignalValue, uint32 SignalLength);
typedef uint8 {MODULE_NAME}_TxCanSigalPropsType;

typedef struct
{{
    uint16 SignalLengthInByte;
    uint16 SignalBufferStartIdx;
    uint16 TxSetSignalIdx;
    {MODULE_NAME}_TxCanSigalPropsType TxSigalProps;
}} {MODULE_NAME}_TxCanSignalConfigType;

/********************************************************************************************************************
 *  PUBLIC CONSTANT AND VARIABLE DECLARATION
 ********************************************************************************************************************/
extern const {MODULE_NAME}_TxCanSignalConfigType {MODULE_NAME}_TxCanSignalConfig[{ENUM_END}];
extern const {MODULE_NAME}_TxCanSetSignalFunctPtrType {MODULE_NAME}_TxCanSetSignalFunctPtrConfig[CAN{PREFIX}_NUM_OF_SET_SIG_FUNCT];
extern const {MODULE_NAME}_TxCanSetSignalU8ArrFunctPtrType {MODULE_NAME}_TxCanSetSignalU8ArrFunctPtrConfig[CAN{PREFIX}_NUM_OF_SET_SIG_U8ARR_FUNCT];

/********************************************************************************************************************
 *  PUBLIC FUNCTION PROTOTYPE
 ********************************************************************************************************************/
''')

        for name in analysis['set_sig_list']:
            f.write(f'extern Std_ReturnType {MODULE_NAME}_TxCanSetSig_{name}(uint8 *pSignalValue);\n')
        for name in analysis['set_sig_list_u8Arr']:
            f.write(f'extern Std_ReturnType {MODULE_NAME}_TxCanSetSig_{name}(uint8 *pSignalValue, uint32 SignalLength);\n')

        if analysis['general_iso_support']:
            f.write(f'''
#ifdef {PREFIX}_TXCAN_ISO_SIGNAL_SUPPORTED
/*!******************************************************************************************************************
*   @brief   This method is used to transmit the Iso signal value to the COM.
*   @param[in] SignalId TX signal id
*   @param[in] pTxSignalData pointer to the data to be updated in COM
*   @param[in] SignalLength length of the data to be udpated
*   @return Std_ReturnType status of the request
*   E_OK status accepted
*   E_NOT_OK status not accepted
********************************************************************************************************************/
extern Std_ReturnType CComAbsMdl_TxCan_SetIsoSignal(tSafeTxSignalId SignalId, uint8 *pTxSignalData, uint32 SignalLength);
/*!******************************************************************************************************************
*   @brief   This method is called from ComAbsMdl_Tx on successful transmission of the signal CanTp signal.
*   Based on the last transmission request the corresponding ISO signal confirmation to be udpated.
********************************************************************************************************************/
extern void CComAbsMdl_TxCan_SetIsoSignalConfirmation(tSafeTxSignalId SignalId);
#endif
''')

        f.write(f'''
/********************************************************************************************************************
 *  Below are the set of signals that can be send via CComAbsMdl_TxCan_SetSignalValue API
 ********************************************************************************************************************/
/* SignalEnum SignalEnumValue SignalNameInDbc MessageName
''')
        for i, sig in enumerate(signals):
            if sig['LengthInBytes'] <= 4:
                f.write(f"   {sig['VSignalEnum']}    {i}     {sig['SignalName']}    {sig['MsgName']}\n")
        f.write("*/\n")

        f.write(f'''
/********************************************************************************************************************
 *  Below are the set of signals that can be send via CComAbsMdl_TxCan_SetSignalValueU8Arr API
 ********************************************************************************************************************/
/* SignalEnum SignalEnumValue SignalNameInDbc MessageName
''')
        for i, sig in enumerate(signals):
            if sig['LengthInBytes'] > 4:
                f.write(f"   {sig['VSignalEnum']}    {i}     {sig['SignalName']}    {sig['MsgName']}\n")
        f.write("*/\n")

        f.write(f'''
#define {PREFIX}_SEC_CODE_STOP
#define {PREFIX}_CORE_CONST_SEC_END
#define {PREFIX}_CORE_DATA_SEC_END
#define {PREFIX}_CORE_BSS_SEC_END
#include "MemMap.h"
#endif /* {PREFIX}_TX_CFG_H_ */
/*******************************************************************************************************************
 *                                 END OF FILE
 ********************************************************************************************************************/
''')

def generate_tx_source_file(signals: List[Dict[str, Any]], analysis: Dict[str, Any], user_name: str, gen_time: str, output_dir: str):
    source_path = os.path.join("..\gen-files", SOURCE_FILE)
    os.makedirs(output_dir, exist_ok=True)
    with open(source_path, "w") as f:
        f.write(f'''/********************************************************************************************************************
 * C O P Y R I G H T
 ********************************************************************************************************************
 * Copyright (c) 2019 by Visteon Corporation.       All rights reserved.
 *
 * NOTICE: This Is an unpublished work Of authorship, which contains trade secrets.
 * Visteon Corporation owns all rights to this work And intends to maintain it in confidence to preserve
 * its trade secret status. Visteon Corporation reserves the right, under the copyright laws of the United States
 * Or those of any other country that may have jurisdiction, to protect this work as an unpublished work,
 * in the event of an inadvertent Or deliberate unauthorized publication. Visteon Corporation also reserves its
 * rights under all copyright laws to protect this work as a published work, when appropriate.
 * Those having access to this work may Not copy it, use it, modify it, Or disclose the information contained in it
 * without the written authorization of Visteon Corporation.
 ********************************************************************************************************************/
/*!*****************************************************************************************************************
 *    @file {SOURCE_FILE}
 *    @ingroup {MODULE_NAME}
 *    @brief This file contains the configuration of the {MODULE_NAME} component.
 *    @Note : Auto Generated code .Do Not hand modify this file
 *    Auto generated code in {user_name} at {gen_time}
 ********************************************************************************************************************/
#include "ComAbsMdlSafe.h"
#include "{HEADER_FILE}"
#include "Rte_CComAbsMdlSafe_Type.h"
#include "Rte_Type.h"

#define {PREFIX}_SEC_CODE_START
#define {PREFIX}_CORE_CONST_SEC_START
#define {PREFIX}_CORE_DATA_SEC_START
#define {PREFIX}_CORE_BSS_SEC_START
#include "MemMap.h"

/********************************************************************************************************************
 *  PUBLIC CONSTANT AND VARIABLE DEFINITION
 ********************************************************************************************************************/
const {MODULE_NAME}_TxCanSetSignalFunctPtrType {MODULE_NAME}_TxCanSetSignalFunctPtrConfig[CAN{PREFIX}_NUM_OF_SET_SIG_FUNCT] =
{{\n''')

        for i, name in enumerate(analysis['set_sig_list']):
            comma = "" if i == len(analysis['set_sig_list']) - 1 else ","
            f.write(f"    &{MODULE_NAME}_TxCanSetSig_{name}{comma}  /* {i} - {analysis['set_sig_list_props'][name]} */\n")
        f.write("};\n\n")

        f.write(f"const {MODULE_NAME}_TxCanSetSignalU8ArrFunctPtrType {MODULE_NAME}_TxCanSetSignalU8ArrFunctPtrConfig[CAN{PREFIX}_NUM_OF_SET_SIG_U8ARR_FUNCT] =\n{{\n")
        if analysis['set_sigArr'] > 0:
            for i, name in enumerate(analysis['set_sig_list_u8Arr']):
                comma = "" if i == len(analysis['set_sig_list_u8Arr']) - 1 else ","
                f.write(f"    &{MODULE_NAME}_TxCanSetSig_{name}{comma}  /* {i} - {analysis['set_sig_list_u8Arr_props'][name]} */\n")
        else:
            f.write("    NULL_PTR\n")
        f.write("};\n\n")

        f.write(f"const {MODULE_NAME}_TxCanSignalConfigType {MODULE_NAME}_TxCanSignalConfig[{ENUM_END}] =\n{{\n")
        l_u8 = l_u16 = l_u32 = l_u64 = 0
        set_id = set_id_arr = 0
        for i, sig in enumerate(signals):
            bits = sig['Lengthinbits']
            bytes_len = sig['LengthInBytes']
            buf_idx = 0

            if bits <= 8 and sig['ValueRetention']:
                buf_idx = l_u8
                l_u8 += 1
            elif bits <= 16 and sig['ValueRetention']:
                buf_idx = l_u16
                l_u16 += 1
            elif bits <= 32 and sig['ValueRetention']:
                buf_idx = l_u32
                l_u32 += 1
            elif sig['ValueRetention']:
                buf_idx = l_u64
                l_u64 += bytes_len
            else:
                if bits <= 8:
                    buf_idx = f"{PREFIX}_TXCAN_NUM_OF_U8_BUFFERS"
                elif bits <= 16:
                    buf_idx = f"{PREFIX}_TXCAN_NUM_OF_U16_BUFFERS"
                elif bits <= 32:
                    buf_idx = f"{PREFIX}_TXCAN_NUM_OF_U32_BUFFERS"
                else:
                    buf_idx = f"{PREFIX}_TXCAN_NUM_OF_U8ARR_BUFFERS"

            set_idx = set_id if bytes_len <= 4 else set_id_arr
            if bytes_len <= 4:
                set_id += 1
            else:
                set_id_arr += 1

            flags = [f"{PREFIX}_TXCAN_CONFIRMATION_ENABLED"]
            if sig['ValueRetention']:
                flags.append(f"{PREFIX}_TXCAN_KEEP_ALIVE_BIT_ENABLED")
            if sig['SignalIsoSupport']:
                flags.append(f"{PREFIX}_TXCAN_ISO_SUPPORTED")
            flag_str = " | ".join(flags)

            comma = "" if i == len(signals) - 1 else ","
            f.write(f"    {{ {bytes_len}U, {buf_idx}, {set_idx}U, ({flag_str}) }}{comma} /* {i} {sig['SignalName']} */\n")
        f.write("};\n\n")

        f.write(f"#define {PREFIX}_STOP_SEC_CONST\n#include \"MemMap.h\"\n\n")
        f.write(f"/* === Place all function code into Safe/ASIL CODE section === */\n")
        f.write(f"#define {PREFIX}_START_SEC_CODE\n#include \"MemMap.h\"\n\n")
        for sig in signals:
            name = sig['SignalName']
        if sig['InternalSignal'] == False:
            if sig['LengthInBytes'] <= 4:           
                f.write(f'''/********************************************************************************************************************
*   @brief   This method Is used to send the signal {name} to Com
*   @param[in] pSignalValue pointer to signal data
********************************************************************************************************************/
Std_ReturnType {MODULE_NAME}_TxCanSetSig_{name}(uint8 *pSignalValue)
{{
    {sig['MsgName']} msg;
    Rte_TransformerError transformerError;
    Std_ReturnType ret = E_NOT_OK;
    if (pSignalValue == NULL_PTR) {{ return E_NOT_OK; }}

    msg.{name} = *pSignalValue;
    ret = Rte_Write_TrustecSWC_PP_TX_{sig['MsgName']}_{sig['MsgName']}(&msg, &transformerError);
    return ret;
}}\n\n''')
            else:
                f.write(f'''/********************************************************************************************************************
*   @brief   This method Is used to send the signal {name} to Com
*   @param[in] pSignalValue pointer to signal data
*   @param[in] SignalLength signal length
********************************************************************************************************************/
Std_ReturnType {MODULE_NAME}_TxCanSetSig_{name}(uint8 *pSignalValue, uint32 SignalLength)
{{
    if (pSignalValue == NULL_PTR) {{ return E_NOT_OK; }}
    {sig['MsgName']} msg;
    Rte_TransformerError transformerError;
    Std_ReturnType ret = E_NOT_OK;
    memcpy(msg.{name}, pSignalValue, SignalLength);
    ret = Rte_Write_TrustecSWC_PP_TX_{sig['MsgName']}_SG_{sig['MsgName']}(&msg, &transformerError);
    return ret;
}}\n\n''')

        f.write(f'''/********************************************************************************************************************
*   @brief   Generic function to route signal to correct setter
*   @param[in] id Signal ID
*   @param[in] pSignalValue pointer to signal data
*   @return Std_ReturnType status
********************************************************************************************************************/
Std_ReturnType {MODULE_NAME}_TxCanPutTo_COM_HS_Core(tSafeTxSignalId id, const uint8 *pSignalValue)
{{
    if (pSignalValue == NULL_PTR) {{ return E_NOT_OK; }}
    switch (id)
    {{\n''')
        for sig in signals:
            f.write(f"        case {sig['VSignalEnum']}:\n")
            f.write(f"            return {MODULE_NAME}_TxCanSetSig_{sig['SignalName']}((uint8*)pSignalValue);\n")
        f.write(f"        default: return E_NOT_OK;\n    }}\n}}\n")

        f.write(f'''
#define {PREFIX}_SEC_CODE_STOP
#define {PREFIX}_CORE_CONST_SEC_END
#define {PREFIX}_CORE_DATA_SEC_END
#define {PREFIX}_CORE_BSS_SEC_END
#include "MemMap.h"
/*******************************************************************************************************************
 *                                 END OF FILE
 ********************************************************************************************************************/
''')

def generate_code_comabs_tx_cfg(excel_name: str, user_name: str, gen_time: str, output_dir: str = "."):
    signals = load_tx_signals(excel_name)
    analysis = analyze_tx_signals(signals)
    generate_tx_header_file(signals, analysis, user_name, gen_time, output_dir)
    generate_tx_source_file(signals, analysis, user_name, gen_time, output_dir)